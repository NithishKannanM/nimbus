from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

W, H = 13.333, 7.5
INK    = RGBColor(0x0F, 0x17, 0x2A)
BODY   = RGBColor(0x33, 0x41, 0x55)
MUTED  = RGBColor(0x64, 0x74, 0x8B)
ACCENT = RGBColor(0x0F, 0x76, 0x6E)
ACCSOF = RGBColor(0xCC, 0xFB, 0xF1)
RULE   = RGBColor(0xE2, 0xE8, 0xF0)
PANEL  = RGBColor(0xF8, 0xFA, 0xFC)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
AMBER  = RGBColor(0xB4, 0x53, 0x09)
FONT   = "Segoe UI"
MONO   = "Consolas"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    return sh


def tb(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, text, size=14, color=BODY, bold=False, font=FONT, space_before=0,
         space_after=6, align=PP_ALIGN.LEFT, first=False, line=1.25, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color; r.font.name = font
    return p


def rich(tf, chunks, size=14, space_before=0, space_after=6, first=False, line=1.25,
         align=PP_ALIGN.LEFT):
    """chunks = [(text, bold, color, font_or_None)]"""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before); p.space_after = Pt(space_after)
    p.line_spacing = line
    for text, bold, color, fnt in chunks:
        r = p.add_run(); r.text = text
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color; r.font.name = fnt or FONT
    return p


def header(s, kicker, title):
    rect(s, 0, 0, W, 0.055, fill=ACCENT)
    t = tb(s, 0.85, 0.42, 11.6, 0.9)
    para(t, kicker.upper(), size=10.5, color=ACCENT, bold=True, space_after=3, first=True)
    para(t, title, size=29, color=INK, bold=True, space_after=0, line=1.05)
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.85), Inches(1.52),
                                Inches(12.48), Inches(1.52))
    ln.line.color.rgb = RULE; ln.line.width = Pt(1.0)


def footer(s, n):
    t = tb(s, 0.85, 6.92, 11.6, 0.3)
    p = para(t, f"Nimbus  ·  Review 1  ·  {n}", size=9, color=MUTED, space_after=0, first=True)
    p.alignment = PP_ALIGN.RIGHT


def arrowhead(conn):
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    ln.append(tail)


def arrow(s, x1, y1, x2, y2, color=MUTED, lw=1.25, head=True):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                               Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(lw)
    if head:
        arrowhead(c)
    return c


def boxtext(sh, lines):
    """lines = [(text, size, bold, color, font_or_None, space_after)]"""
    tf = sh.text_frame
    for i, (text, size, bold, color, fnt, sa) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(0); p.space_after = Pt(sa); p.line_spacing = 1.05
        r = p.add_run(); r.text = text
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color; r.font.name = fnt or FONT
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE


# ─────────────────────────────────────────────── 1. TITLE
s = slide()
rect(s, 0, 0, W, H, fill=WHITE)
rect(s, 0, 0, 0.28, H, fill=ACCENT)
rect(s, 0.28, 0, 0.06, H, fill=ACCSOF)

t = tb(s, 1.35, 1.95, 10.6, 1.2)
para(t, "NIMBUS", size=15, color=ACCENT, bold=True, space_after=10, first=True)
para(t, "Telemetry-Aware Adaptive", size=44, color=INK, bold=True, space_after=0, line=1.02)
para(t, "Workload Scheduling on AWS", size=44, color=INK, bold=True, space_after=0, line=1.02)

ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.35), Inches(4.28),
                            Inches(4.35), Inches(4.28))
ln.line.color.rgb = ACCENT; ln.line.width = Pt(2.5)

t = tb(s, 1.35, 4.52, 9.9, 0.9)
para(t, "Can a scheduler that reads live worker telemetry place jobs better than one "
        "that reasons only about declared and reserved capacity?",
     size=15.5, color=BODY, space_after=0, line=1.35, first=True)

b = rect(s, 1.35, 5.72, 2.15, 0.36, fill=ACCSOF, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.4)
boxtext(b, [("REVIEW 1", 10.5, True, ACCENT, None, 0)])

t = tb(s, 3.75, 5.74, 8.2, 0.34)
para(t, "Phase 3 of 9 complete  ·  274 tests passing", size=11.5, color=MUTED,
     space_after=0, first=True)

t = tb(s, 1.35, 6.42, 10.6, 0.7)
rich(t, [("Presented by  ", False, MUTED, None), ("〈 Name 〉", True, INK, None),
         ("   ·   Reg. No.  ", False, MUTED, None), ("〈 Number 〉", True, INK, None)],
     size=12, space_after=4, first=True)
rich(t, [("Guide  ", False, MUTED, None), ("〈 Guide Name 〉", True, INK, None),
         ("   ·   Department of 〈 … 〉", False, MUTED, None)], size=12, space_after=0)

# ─────────────────────────────────────────────── 2. PROBLEM STATEMENT
s = slide()
header(s, "Problem Statement", "Schedulers place work using numbers nobody measured")
footer(s, "02")

t = tb(s, 0.85, 1.85, 7.35, 3.6)
para(t, "A cluster scheduler must answer one question for every job: which machine "
        "should run this?", size=15, color=INK, bold=True, space_after=10, line=1.3, first=True)
para(t, "Production placement engines answer it from paperwork rather than from "
        "measurement:", size=14, color=BODY, space_after=10, line=1.35)

for head, tail in [
    ("Declared requirements", "what the job's author claimed it would need"),
    ("Reserved capacity", "the arithmetic sum of those claims on each host"),
    ("Static attributes", "instance type, AZ, labels — fixed at registration"),
]:
    rich(t, [("▪  ", False, ACCENT, None), (head + " — ", True, INK, None),
             (tail, False, BODY, None)], size=13.5, space_after=6, line=1.3)

para(t, "None of these is the machine's actual condition.", size=14.5, color=INK,
     bold=True, space_before=8, space_after=7, line=1.3)
para(t, "A host can sit at 100% CPU from work the scheduler never dispatched — a "
        "co-tenant, a backup job, an operator's build. Nothing was reserved, so it "
        "still advertises full capacity as free.",
     size=13, color=BODY, space_after=0, line=1.3)

p = rect(s, 8.65, 1.85, 3.83, 3.6, fill=PANEL, line=RULE)
tf = p.text_frame
tf.margin_left = tf.margin_right = Inches(0.28)
tf.margin_top = Inches(0.24)
tf.vertical_anchor = MSO_ANCHOR.TOP
para(tf, "THE BLIND SPOT, CONCRETELY", size=10, color=ACCENT, bold=True,
     space_after=12, first=True)
rich(tf, [("Amazon ECS", True, INK, None)], size=13, space_after=4)
para(tf, "Its cluster query language has no subject for observed CPU or memory "
         "utilisation. Placement filters and strategies read declared and reserved "
         "values only.", size=11.5, color=BODY, space_after=12, line=1.32)
rich(tf, [("AWS Batch", True, INK, None)], size=13, space_after=4)
para(tf, "Fair-share scheduling reasons about consumed queue share per tenant — "
         "not about host pressure.", size=11.5, color=BODY, space_after=0, line=1.32)

r = rect(s, 0.85, 6.06, 11.63, 0.74, fill=WHITE, line=ACCENT, lw=1.25)
tf = r.text_frame
tf.margin_left = tf.margin_right = Inches(0.26); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
rich(tf, [("The gap.  ", True, ACCENT, None),
          ("Does adding observed, high-frequency, per-host telemetry to the placement "
           "function measurably improve scheduling — or does reservation accounting "
           "already capture everything that matters?", False, INK, None)],
     size=13.5, space_after=0, line=1.3, first=True)

# ─────────────────────────────────────────────── 3. WHAT THIS IS NOT
s = slide()
header(s, "Problem Statement", "Scoping: what this project does not claim")
footer(s, "03")

t = tb(s, 0.85, 1.82, 11.63, 0.6)
para(t, "AWS already solves production job scheduling. Nimbus is not a replacement "
        "for any of it, and the project is forbidden from claiming otherwise.",
     size=14, color=BODY, space_after=0, line=1.35, first=True)

cols = [
    ("ALREADY SOLVED BY AWS", MUTED, RULE, PANEL, [
        "FIFO queues, job priority, dependsOn DAGs",
        "Fair-share scheduling with usage decay",
        "binpack / spread / random placement",
        "Lease semantics, retries, DLQs",
        "Target-tracking and predictive autoscaling",
        "Capacity-aware instance allocation",
    ]),
    ("DELIBERATELY OUT OF SCOPE", AMBER, RULE, PANEL, [
        "Multi-tenant fairness — starvation control only",
        "Autoscaling — the fleet is fixed at 3 workers",
        "Instance-type selection",
        "DAG dependencies between jobs",
        "Priority ordering — designed, not implemented",
        "Production security posture",
    ]),
    ("THE NARROW THING NIMBUS TESTS", ACCENT, ACCENT, WHITE, [
        "Observed per-host CPU and memory pressure",
        "Smoothed over a short sample window",
        "Plus a least-squares utilisation trend",
        "Added to the per-job placement score",
        "Measured against the static baseline",
        "Treated as a hypothesis that may be rejected",
    ]),
]

x = 0.85
for title, tcol, border, bg, items in cols:
    box = rect(s, x, 2.62, 3.71, 3.62, fill=bg, line=border, lw=1.25)
    tf = box.text_frame
    tf.margin_left = tf.margin_right = Inches(0.24)
    tf.margin_top = Inches(0.22)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    para(tf, title, size=10, color=tcol, bold=True, space_after=13, first=True, line=1.15)
    for it in items:
        rich(tf, [("·  ", False, tcol, None), (it, False, BODY, None)],
             size=11.5, space_after=8, line=1.28)
    x += 3.96

tf = tb(s, 0.85, 6.46, 11.63, 0.36)
para(tf, "The real test is not “does Adaptive beat FIFO” — it should, almost "
         "arithmetically. It is “does Adaptive beat Static Resource-Aware”.",
     size=12.5, color=INK, italic=True, space_after=0, first=True)

# ─────────────────────────────────────────────── 4. OBJECTIVE
s = slide()
header(s, "Objective", "One hypothesis, three falsifiable sub-claims")
footer(s, "04")

q = rect(s, 0.85, 1.82, 11.63, 1.06, fill=ACCSOF, line=None)
tf = q.text_frame
tf.margin_left = Inches(0.3); tf.margin_right = Inches(0.3)
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
rich(tf, [("H1.  ", True, ACCENT, None),
          ("For heterogeneous workloads on workers with dynamically changing resource "
           "utilisation, placement decisions informed by real-time telemetry and workload "
           "characteristics may improve completion latency, utilisation and load balance "
           "compared with FIFO or static resource-aware scheduling.", False, INK, None)],
     size=13.5, space_after=0, line=1.32, first=True)
rect(s, 0.85, 1.82, 0.055, 1.06, fill=ACCENT)

hdrs = ["ID", "Sub-hypothesis", "Primary metric", "Null"]
widths = [0.85, 5.65, 2.9, 2.23]
rowdata = [
    ("H1a", "Adaptive reduces tail completion latency vs. FIFO", "p95 completion time", "no difference"),
    ("H1b", "Adaptive improves load balance vs. FIFO", "CV of worker busy-time", "no difference"),
    ("H1c", "Adaptive beats Static Resource-Aware — live telemetry adds\n"
            "something over reservation accounting", "p95 completion time", "no difference"),
]

y = 3.12
x = 0.85
for wdt, h in zip(widths, hdrs):
    c = rect(s, x, y, wdt, 0.4, fill=INK, line=None)
    tf = c.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.14)
    para(tf, h, size=10.5, color=WHITE, bold=True, space_after=0, first=True)
    x += wdt
y += 0.4

for i, row in enumerate(rowdata):
    hgt = 0.72 if i == 2 else 0.5
    x = 0.85
    bg = WHITE if i % 2 == 0 else PANEL
    for j, (wdt, cell) in enumerate(zip(widths, row)):
        c = rect(s, x, y, wdt, hgt, fill=bg, line=RULE, lw=0.75)
        tf = c.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.14)
        emph = (i == 2)
        para(tf, cell, size=11.5, color=INK if (j == 0 or emph) else BODY,
             bold=(j == 0), space_after=0, line=1.25, first=True)
        x += wdt
    y += hgt

t = tb(s, 0.85, 5.36, 11.63, 0.4)
rich(t, [("H1c is the real test.  ", True, ACCENT, None),
         ("H1a and H1b compare against baselines that ignore resources entirely. If "
          "Adaptive beats FIFO but not Resource-Aware, the honest conclusion — resource "
          "awareness helps, live telemetry adds nothing measurable — is a result, "
          "not a failure.", False, BODY, None)],
     size=12.5, space_after=0, line=1.3, first=True)

y = 5.98
items = [
    ("Pre-registered", "Weights, metrics, α = 0.05 and the ≥5% materiality bar fixed in writing before the code existed"),
    ("Controlled", "Every strategy uses FIFO ordering, so any difference is attributable to placement alone"),
    ("Replayable", "Every decision is a trace that can be re-scored counterfactually under another policy"),
]
x = 0.85
for head, tail in items:
    box = rect(s, x, y, 3.71, 0.9, fill=PANEL, line=RULE, lw=0.75)
    tf = box.text_frame
    tf.margin_left = tf.margin_right = Inches(0.18); tf.margin_top = Inches(0.1)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, head.upper(), size=9.5, color=ACCENT, bold=True, space_after=4, first=True)
    para(tf, tail, size=10, color=BODY, space_after=0, line=1.22)
    x += 3.96

# ─────────────────────────────────────────────── 5. ARCHITECTURE — OVERVIEW
s = slide()
header(s, "System Architecture Design", "Control plane decides, data plane executes")
footer(s, "05")

def dbox(x, y, w, h, lines, fill=WHITE, line=RULE, lw=1.0, radius=None):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    b = rect(s, x, y, w, h, fill=fill, line=line, lw=lw, shape=shape, radius=radius)
    boxtext(b, lines)
    return b

L = lambda txt, sz=11, bold=False, col=INK, fnt=None, sa=1: (txt, sz, bold, col, fnt, sa)

# client
dbox(5.32, 1.74, 2.7, 0.4, [L("CLIENT   ·   CLI / dashboard", 10.5, False, MUTED)],
     fill=PANEL, line=RULE, radius=0.5)
arrow(s, 6.67, 2.14, 6.67, 2.42)

# api
dbox(4.32, 2.42, 4.7, 0.5,
     [L("REST API   ·   FastAPI", 12, True, INK),
      L("/jobs   /workers   /scheduler   /scheduler/decisions", 9.5, False, MUTED)],
     fill=WHITE, line=INK, lw=1.5)

arrow(s, 5.4, 2.92, 3.95, 3.24)
arrow(s, 7.95, 2.92, 9.4, 3.24)

# stores
dbox(2.35, 3.24, 3.2, 0.66,
     [L("StateStore", 11.5, True, INK),
      L("job metadata  ·  decision traces  ·  telemetry", 9, False, MUTED)],
     fill=PANEL, line=RULE)
dbox(7.8, 3.24, 3.2, 0.66,
     [L("JobQueue", 11.5, True, INK),
      L("lease  ·  ack / nack  ·  defer  ·  DLQ", 9, False, MUTED)],
     fill=PANEL, line=RULE)

arrow(s, 3.95, 3.90, 4.9, 4.22)
arrow(s, 9.4, 3.90, 8.45, 4.22)

# scheduler
sch = rect(s, 2.35, 4.22, 8.65, 1.12, fill=ACCSOF, line=ACCENT, lw=1.5)
tf = sch.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = tf.margin_right = Inches(0.2)
para(tf, "SCHEDULER LOOP   ·   singleton, control plane", size=11.5, color=ACCENT,
     bold=True, align=PP_ALIGN.CENTER, space_after=6, first=True)
rich(tf, [("1 ", True, ACCENT, MONO), ("lease job    ", False, INK, None),
          ("2 ", True, ACCENT, MONO), ("snapshot fleet    ", False, INK, None),
          ("3 ", True, ACCENT, MONO), ("filter infeasible    ", False, INK, None),
          ("4 ", True, ACCENT, MONO), ("score candidates    ", False, INK, None),
          ("5 ", True, ACCENT, MONO), ("persist trace    ", False, INK, None),
          ("6 ", True, ACCENT, MONO), ("dispatch", False, INK, None)],
     size=10.5, space_after=0, align=PP_ALIGN.CENTER, line=1.1)

t = tb(s, 11.15, 4.42, 1.35, 0.75)
para(t, "push\ndispatch —\nworkers never\npull", size=9, color=MUTED, space_after=0,
     line=1.2, first=True)

for wx in (3.5, 6.24, 8.98):
    arrow(s, 6.67, 5.34, wx + 1.1, 5.72, lw=1.0)

# workers
for i, wx in enumerate((3.5, 6.24, 8.98), start=1):
    dbox(wx, 5.72, 2.2, 0.78,
         [L(f"WORKER {i}", 10.5, True, INK),
          L("executor  ·  telemetry  ·  heartbeat", 8.5, False, MUTED)],
         fill=WHITE, line=INK, lw=1.0)

for wx in (3.5, 6.24, 8.98):
    a = arrow(s, wx + 1.1, 6.50, 6.67, 6.74, color=ACCENT, lw=1.0)

t = tb(s, 7.0, 6.62, 5.5, 0.3)
para(t, "telemetry + heartbeat + result  →  StateStore", size=9.5, color=ACCENT,
     space_after=0, first=True)

t = tb(s, 0.85, 5.74, 2.5, 0.8)
para(t, "DATA PLANE", size=9.5, color=MUTED, bold=True, space_after=3, first=True)
para(t, "Accepts, executes, reports. Never chooses what to run.",
     size=9.5, color=MUTED, space_after=0, line=1.25)

# ─────────────────────────────────────────────── 6. ARCHITECTURE — INFORMATION BOUNDARY
s = slide()
header(s, "System Architecture Design", "The information boundary is the experiment")
footer(s, "06")

t = tb(s, 0.85, 1.82, 11.63, 0.5)
para(t, "Worker state is split into three categories that are never conflated, because "
        "the central experiment depends on telling them apart. A placement policy never "
        "receives raw worker state — it declares its information budget through a view.",
     size=13.5, color=BODY, space_after=0, line=1.35, first=True)

specs = [
    ("DECLARED", "Operator configuration", "vCPU, memory and slot capacity set at\nregistration. Never changes at runtime.",
     "read by every policy", MUTED, PANEL, RULE),
    ("RESERVED", "Nimbus bookkeeping", "Sum of what Nimbus has dispatched and\nnot yet reclaimed. reserve() / release().",
     "read by every policy", MUTED, PANEL, RULE),
    ("OBSERVED", "psutil, sampled live", "Smoothed CPU mean, memory pressure and\nleast-squares trend over the last N samples.",
     "read by adaptive-v1 only", ACCENT, WHITE, ACCENT),
]

x = 0.85
for name, src, desc, who, col, bg, border in specs:
    box = rect(s, x, 2.56, 3.71, 2.06, fill=bg, line=border, lw=1.5 if col == ACCENT else 1.0)
    tf = box.text_frame
    tf.margin_left = tf.margin_right = Inches(0.24); tf.margin_top = Inches(0.2)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    para(tf, name, size=15, color=col if col == ACCENT else INK, bold=True,
         space_after=2, first=True)
    para(tf, src, size=10, color=col, space_after=10)
    para(tf, desc, size=11, color=BODY, space_after=10, line=1.28)
    rich(tf, [("→  ", False, col, None), (who, True, col, None)], size=10.5, space_after=0)
    x += 3.96

t = tb(s, 0.85, 4.74, 11.63, 0.35)
para(t, "Two views, pinned by tests, enforce that boundary:", size=13, color=INK,
     bold=True, space_after=0, first=True)

vrows = [
    ("StaticWorkerView", "round_robin, static_resource_aware, static_resource_aware_binpack",
     "DECLARED + RESERVED", MUTED),
    ("AdaptiveWorkerView", "adaptive  (adaptive-v1)",
     "DECLARED + RESERVED + OBSERVED + telemetry quality", ACCENT),
]
y = 5.10
for vname, users, sees, col in vrows:
    b = rect(s, 0.85, y, 11.63, 0.50, fill=WHITE, line=RULE, lw=0.75)
    tf = b.text_frame
    tf.margin_left = Inches(0.22); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    rich(tf, [(f"{vname:<22}", True, col, MONO), (f"  {users}", False, BODY, None),
              ("      sees:  ", False, MUTED, None), (sees, True, col, None)],
         size=11, space_after=0, first=True)
    y += 0.56

r = rect(s, 0.85, 6.26, 11.63, 0.56, fill=PANEL, line=None)
tf = r.text_frame
tf.margin_left = tf.margin_right = Inches(0.22)
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
rich(tf, [("Recorded ≠ consumed.  ", True, INK, None),
          ("Traces record all three categories for every candidate, written after the "
           "policy has returned — so replay can prove which fields a decision actually "
           "depended on.", False, BODY, None)],
     size=11, space_after=0, first=True, line=1.25)

# ─────────────────────────────────────────────── 7. ARCHITECTURE — PLACEMENT
s = slide()
header(s, "System Architecture Design", "Ordering and placement are separate objects")
footer(s, "07")

t = tb(s, 0.85, 1.80, 11.63, 0.36)
para(t, "Every strategy uses FIFO ordering, so a comparison isolates placement — "
        "the only variable is what the policy may see.",
     size=12.5, color=BODY, space_after=0, line=1.3, first=True)

hdrs = ["Strategy", "Placement", "Sees", "Status"]
widths = [3.05, 4.15, 2.85, 1.58]
rows = [
    ("fifo", "round-robin rotation", "DECLARED + RESERVED", "implemented", MUTED),
    ("static_resource_aware", "balance on projected utilisation", "DECLARED + RESERVED", "implemented", MUTED),
    ("static_resource_aware_binpack", "ECS-style binpack", "DECLARED + RESERVED", "implemented", MUTED),
    ("adaptive", "adaptive-v1 weighted score", "+ OBSERVED", "implemented", ACCENT),
    ("priority", "priority + aging", "—", "not implemented", RGBColor(0x94, 0xA3, 0xB8)),
]

y = 2.36
x = 0.85
for wdt, h in zip(widths, hdrs):
    c = rect(s, x, y, wdt, 0.34, fill=INK, line=None)
    tf = c.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.margin_left = Inches(0.14)
    para(tf, h, size=10, color=WHITE, bold=True, space_after=0, first=True)
    x += wdt
y += 0.34

for i, (name, place, sees, status, col) in enumerate(rows):
    x = 0.85
    hot = (col == ACCENT)
    bg = ACCSOF if hot else (WHITE if i % 2 == 0 else PANEL)
    for j, (wdt, cell) in enumerate(zip(widths, (name, place, sees, status))):
        c = rect(s, x, y, wdt, 0.38, fill=bg, line=RULE, lw=0.75)
        tf = c.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.margin_left = Inches(0.14)
        para(tf, cell, size=10.5,
             color=col if (j == 0 or j == 3 or hot) else BODY,
             bold=(j == 0 or hot), font=MONO if j == 0 else FONT,
             italic=(status == "not implemented"), space_after=0, first=True)
        x += wdt
    y += 0.38

t = tb(s, 0.85, 4.76, 5.7, 0.34)
para(t, "adaptive-v1 — score, frozen before the code was written", size=12.5,
     color=INK, bold=True, space_after=0, first=True)

f = rect(s, 0.85, 5.08, 5.7, 1.28, fill=INK, line=None)
tf = f.text_frame
tf.margin_left = Inches(0.24); tf.margin_top = Inches(0.16)
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
for i, (sym, term, note) in enumerate([
    ("α ·", "S_cpu", "observed, smoothed"),
    ("β ·", "S_memory", "observed"),
    ("γ ·", "S_reservation", "the static model, kept"),
    ("δ ·", "S_queue", "declared + reserved"),
    ("ε ·", "S_trend", "least-squares slope"),
]):
    lead = "Score(w,j) = " if i == 0 else "             + "
    rich(tf, [(lead, False, RGBColor(0x94, 0xA3, 0xB8), MONO),
              (sym + " ", False, ACCSOF, MONO),
              (f"{term:<15}", True, WHITE, MONO),
              (note, False, RGBColor(0x94, 0xA3, 0xB8), None)],
         size=9.5, space_after=2, first=(i == 0), line=1.1)

t = tb(s, 6.9, 4.76, 5.58, 0.34)
para(t, "Weights vary by workload class", size=12.5, color=INK, bold=True,
     space_after=0, first=True)

wt_hdr = ["Class", "α", "β", "γ", "δ", "ε"]
wt_w = [2.08, 0.7, 0.7, 0.7, 0.7, 0.7]
wt_rows = [
    ("CPU_BOUND", "0.50", "0.10", "0.20", "0.15", "0.05", 1),
    ("MEMORY_BOUND", "0.10", "0.50", "0.20", "0.15", "0.05", 2),
    ("IO_BOUND", "0.15", "0.15", "0.25", "0.40", "0.05", 4),
    ("LIGHTWEIGHT", "0.10", "0.05", "0.25", "0.55", "0.05", 4),
]
y = 5.08
x = 6.9
for wdt, h in zip(wt_w, wt_hdr):
    c = rect(s, x, y, wdt, 0.28, fill=INK, line=None)
    tf = c.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, h, size=9.5, color=WHITE, bold=True, space_after=0,
         align=PP_ALIGN.CENTER if h != "Class" else PP_ALIGN.LEFT, first=True)
    if h == "Class":
        tf.margin_left = Inches(0.12)
    x += wdt
y += 0.28
for name, *vals in wt_rows:
    peak = vals[-1]
    vals = vals[:-1]
    x = 6.9
    c = rect(s, x, y, wt_w[0], 0.265, fill=WHITE, line=RULE, lw=0.75)
    tf = c.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.margin_left = Inches(0.12)
    para(tf, name, size=9, color=INK, font=MONO, space_after=0, first=True)
    x += wt_w[0]
    for k, v in enumerate(vals, start=1):
        hot = (k == peak)
        c = rect(s, x, y, wt_w[k], 0.265, fill=ACCSOF if hot else WHITE, line=RULE, lw=0.75)
        tf = c.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tf, v, size=9, color=ACCENT if hot else BODY, bold=hot, font=MONO,
             align=PP_ALIGN.CENTER, space_after=0, first=True)
        x += wt_w[k]
    y += 0.265

t = tb(s, 0.85, 6.44, 10.0, 0.34)
rich(t, [("Telemetry ranks; it never changes eligibility.  ", True, ACCENT, None),
         ("Feasibility is evaluated on static fields alone, stale telemetry falls back to "
          "reservation estimates rather than to “0% CPU”, and there is no threshold cliff "
          "in the score.", False, BODY, None)],
     size=11, space_after=0, line=1.28, first=True)

# ─────────────────────────────────────────────── 8. REFERENCES
s = slide()
header(s, "References", "Sources and prior art")
footer(s, "08")

groups = [
    ("AWS SERVICE DOCUMENTATION — PRIOR ART", 1.78, [
        ("Amazon Web Services, “AWS Batch job scheduling,” AWS Batch User Guide.",
         "docs.aws.amazon.com/batch/latest/userguide/job_scheduling.html"),
        ("Amazon Web Services, “Fair-share scheduling policies,” AWS Batch User Guide.",
         "docs.aws.amazon.com/batch/latest/userguide/fair-share-scheduling.html"),
        ("Amazon Web Services, “Compute environment allocation strategies,” AWS Batch User Guide.",
         "docs.aws.amazon.com/batch/latest/userguide/allocation-strategies.html"),
        ("Amazon Web Services, “Amazon ECS task placement,” Amazon ECS Developer Guide.",
         "docs.aws.amazon.com/AmazonECS/latest/developerguide/task-placement.html"),
        ("Amazon Web Services, “Amazon ECS task placement strategies,” Amazon ECS Developer Guide.",
         "docs.aws.amazon.com/AmazonECS/latest/developerguide/task-placement-strategies.html"),
        ("Amazon Web Services, “Cluster query language,” Amazon ECS Developer Guide.",
         "docs.aws.amazon.com/AmazonECS/latest/developerguide/cluster-query-language.html"),
        ("Amazon Web Services, “Amazon SQS visibility timeout,” Amazon SQS Developer Guide.",
         "docs.aws.amazon.com/AWSSimpleQueueService/latest/SQSDeveloperGuide/sqs-visibility-timeout.html"),
        ("Amazon Web Services, “Amazon SQS dead-letter queues,” Amazon SQS Developer Guide.",
         "docs.aws.amazon.com/AWSSimpleQueueService/latest/SQSDeveloperGuide/sqs-dead-letter-queues.html"),
        ("Amazon Web Services, “Amazon SQS fair queues,” Amazon SQS Developer Guide.",
         "docs.aws.amazon.com/AWSSimpleQueueService/latest/SQSDeveloperGuide/sqs-fair-queues.html"),
        ("Amazon Web Services, “Scale your Auto Scaling group,” Amazon EC2 Auto Scaling User Guide.",
         "docs.aws.amazon.com/autoscaling/ec2/userguide/scale-your-group.html"),
        ("Amazon Web Services, “AWS Price List Query API.”",
         "pricing.us-east-1.amazonaws.com/offers/v1.0/aws/"),
    ]),
]

n = 1
for gtitle, gy, items in groups:
    t = tb(s, 0.85, gy, 11.63, 0.28)
    para(t, gtitle, size=9.5, color=ACCENT, bold=True, space_after=0, first=True)
    y = gy + 0.34
    for cite, url in items:
        t = tb(s, 0.85, y, 11.63, 0.36)
        rich(t, [(f"[{n}]  ", True, MUTED, MONO), (cite, False, INK, None)],
             size=11, space_after=1, first=True, line=1.15)
        para(t, "        " + url, size=9, color=MUTED, font=MONO, space_after=0, line=1.1)
        y += 0.365
        n += 1

t = tb(s, 0.85, 6.22, 5.6, 0.28)
para(t, "TOOLS AND LIBRARIES", size=9.5, color=ACCENT, bold=True, space_after=6, first=True)
rich(t, [("FastAPI  ·  Uvicorn  ·  Pydantic  ·  psutil  ·  NumPy  ·  SciPy  ·  pytest",
          False, BODY, None)], size=10.5, space_after=0)

t = tb(s, 7.0, 6.22, 5.48, 0.28)
para(t, "PROJECT DESIGN DOCUMENTS", size=9.5, color=ACCENT, bold=True, space_after=6, first=True)
rich(t, [("architecture · scheduling · experiments · aws-native-scheduling ·\n"
          "adaptive-v1-preregistration · limitations       (all under docs/)",
          False, BODY, MONO)], size=8.5, space_after=0, line=1.25)

prs.save("Nimbus_Review1.pptx")
print("saved", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
